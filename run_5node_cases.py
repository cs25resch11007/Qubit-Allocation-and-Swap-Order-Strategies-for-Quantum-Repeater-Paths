"""Run path-graph test cases from an Excel workbook in SeQUeNCe with entanglement_swapping.py.

Works with any workbook laid out like 5node.xlsx: one header row, then each test case spans one
row per algorithm, with the case number only on the first row. Columns are found by header name
(see COLUMNS), so their position does not matter. The number of nodes is taken from each case's
p_gen list (n links -> n+1 nodes).

For every test case (first --cases cases of each sheet) and every algorithm row that has a
static swapping order and a qubit allocation, the case is simulated with:

  * p_gen  -> link length (0 < p_gen < 1). Every link uses the same fiber loss, FIBER_LOSS_DB_KM,
             chosen so that p_gen = ANCHOR_PGEN (0.054) gives ANCHOR_KM (100 km):
                FIBER_LOSS_DB_KM = -10 * log10(ANCHOR_PGEN) / ANCHOR_KM  = 0.1268 dB/km
             and each link's length is the one whose transmission equals its p_gen:
                10^(-FIBER_LOSS_DB_KM * L_km / 10) = p_gen   =>   L_km = -10 * log10(p_gen) / FIBER_LOSS_DB_KM
             (0.95 -> 1.8 km, 0.6 -> 17.5 km, 0.3 -> 41.2 km, 0.1 -> 78.9 km, 0.054 -> 100 km).
             With single-heralded generation (default), BSM success 1 and memory/detector
             efficiency 1, an attempt succeeds with probability
                BSM success x (memory eff x detector eff)^2 x transmission = p_gen   (exactly).
             There is no cap on the end-to-end length. With --generation barret_kok the paper's
             Barrett-Kok protocol and 0.95 efficiencies are used instead; attempt success is then
             ~0.41 x p_gen (Barrett-Kok cannot exceed 0.5 per attempt).
  * p_swap -> per-node swapping success probability q_x (Node2..Node{n} -> nodes 1..n-1).
  * Qubit Allocation {node: (left, right)} -> memory pairs per link. Link i joins node i and
             node i+1; its usable pairs are min(right of node i, left of node i+1). Every node
             gets as many memories as the largest left+right in the allocation.
  * Order  -> static swapping order (the last "(a, b, ...)" in the Order cell). Spreadsheet nodes
             are 1-indexed (1..n+1); the simulator uses the paper's 0-indexed nodes (0..n), so
             (4, 3, 2) becomes [3, 2, 1].

Rows without a static order or allocation (e.g. DP, "Dynamic Optimal Co-allocation") cannot be
simulated with a static order; their Expected Entanglement is kept in the output for reference.

Results are appended to <workbook>_results_<generation>.csv (next to the workbook, or --out) after every
run, and finished runs are skipped when the script is started again, so a long batch can be
stopped and resumed.

Usage:
    python local/run_5node_cases.py                                  # 5node.xlsx, all sheets, 10 cases
    python local/run_5node_cases.py --workbook path/to/7node.xlsx    # any workbook, all its sheets
    python local/run_5node_cases.py --workbook my.xlsx --sheets "Wide Random" "Mid-Noise"
    python local/run_5node_cases.py --cases 5 --jobs 4               # first 5 cases, 4 runs in parallel
    python local/run_5node_cases.py --dry-run                        # only print the mapped cases
"""

import argparse
import csv
import math
import os
import re
import sys
import time
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import entanglement_swapping as es

DEFAULT_WORKBOOK = os.path.join(HERE, "5node.xlsx")
CASES_PER_SHEET = 10

# one fiber loss for every link: p_gen = ANCHOR_PGEN maps to ANCHOR_KM
ANCHOR_PGEN = 0.054
ANCHOR_KM = 100
FIBER_LOSS_DB_KM = -10 * math.log10(ANCHOR_PGEN) / ANCHOR_KM   # = 0.1268 dB/km

# entanglement generation used for the links (see the p_gen mapping above)
GENERATION = "single_heralded"      # or "barret_kok" (paper's Section V protocol)
# hardware for single-heralded links, chosen so that attempt success == fiber transmission == p_gen
SH_BSM_SUCCESS_RATE = 1.0
SH_MEMORY_EFFICIENCY = 1.0
SH_DETECTOR_EFFICIENCY = 1.0

# header text of each column used (matched case-insensitively, ignoring surrounding spaces)
COLUMNS = {
    "case": "Test Case",
    "p_gen": "p_gen Chain Configuration",
    "p_swap": "p_swap Node Configuration",
    "algorithm": "Algorithm",
    "allocation": "Qubit Allocation",
    "order": "Order",
    "expected": "Expected Entanglement",
}

# short names for known algorithms; any other algorithm keeps its spreadsheet name
SHORT_NAMES = {
    "Our Algo": "ours",
    "Baseline (CBT, Equal Alloc)": "cbt",
    "(GSO, EA)": "gso",
    "DP": "dp",
}

CSV_FIELDS = ["sheet", "case", "algorithm", "order_sheet", "order_sim", "distances_km", "total_km", "attenuation_db_km",
              "p_gen", "p_swap", "alloc_pairs", "expected_ent_sheet", "reference_expected_ent_sheet",
              "generation", "seed", "e2e", "elapsed_s", "throughput_ent_s", "expired", "wall_s"]


# ----------------------------------------------------------------------------------------
# Reading and mapping the spreadsheet
# ----------------------------------------------------------------------------------------
def parse_values(text):
    """'L1->0.2223, L2->0.3674' -> [0.2223, 0.3674]"""
    return [float(v) for v in re.findall(r"->\s*([0-9.]+)", str(text or ""))]


def parse_order(text):
    """'Greedy Order: (4, 2, 3)' -> (4, 2, 3)  (last tuple in the text), or None if there is none"""
    tuples = re.findall(r"\((\d+(?:\s*,\s*\d+)+)\)", str(text or ""))
    if not tuples:
        return None
    return tuple(int(x) for x in tuples[-1].split(","))


def parse_allocation(text):
    """'{1: (0, 50), 2: (25, 25), ...}' -> {1: (0, 50), 2: (25, 25), ...}"""
    alloc = {}
    for node, left, right in re.findall(r"(\d+)\s*:\s*\(\s*(\d+)\s*,\s*(\d+)\s*\)", str(text or "")):
        alloc[int(node)] = (int(left), int(right))
    return alloc


def pgen_to_km(p):
    """Length whose transmission at FIBER_LOSS_DB_KM equals p (0 < p < 1)."""
    if not 0 < p < 1:
        raise ValueError(f"p_gen must be in (0, 1), got {p}")
    return -10 * math.log10(p) / FIBER_LOSS_DB_KM


def link_pairs(alloc):
    """Usable memory pairs per link: min(right of node i, left of node i+1)."""
    nodes = sorted(alloc)
    return [min(alloc[a][1], alloc[b][0]) for a, b in zip(nodes, nodes[1:])]


def short_name(algorithm):
    algorithm = str(algorithm).strip()
    return SHORT_NAMES.get(algorithm, algorithm)


def column_indices(header, sheet):
    normalized = [str(h).strip().lower() if h is not None else "" for h in header]
    indices = {}
    for key, title in COLUMNS.items():
        if title.lower() not in normalized:
            raise ValueError(f"sheet {sheet!r} has no column {title!r} (found: {[h for h in header if h]})")
        indices[key] = normalized.index(title.lower())
    return indices


def read_cases(workbook, sheets, cases_per_sheet):
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    cases = []
    for sheet in sheets:
        rows = list(wb[sheet].iter_rows(values_only=True))
        col = column_indices(rows[0], sheet)
        sheet_cases = []
        current = None
        for row in rows[1:]:
            def get(key):
                return row[col[key]] if col[key] < len(row) else None

            if get("case") is not None:
                if len(sheet_cases) == cases_per_sheet:
                    break
                current = {"sheet": sheet, "case": int(get("case")), "p_gen": parse_values(get("p_gen")),
                           "p_swap": parse_values(get("p_swap")), "algorithms": {}, "reference": {}}
                if len(current["p_swap"]) != len(current["p_gen"]) - 1:
                    raise ValueError(f"{sheet} case {current['case']}: {len(current['p_gen'])} links need "
                                     f"{len(current['p_gen']) - 1} p_swap values, got {current['p_swap']}")
                sheet_cases.append(current)
            if current is None or get("algorithm") is None:
                continue
            name = short_name(get("algorithm"))
            order, alloc = parse_order(get("order")), parse_allocation(get("allocation"))
            if order is None or not alloc:
                current["reference"][name] = get("expected")   # e.g. DP: no static order to simulate
            else:
                current["algorithms"][name] = {"order_sheet": order, "alloc": alloc, "expected": get("expected")}
        cases.extend(sheet_cases)
    return cases


def build_jobs(cases, seed, generation):
    jobs = []
    for case in cases:
        distances = [pgen_to_km(p) for p in case["p_gen"]]
        attenuations = [FIBER_LOSS_DB_KM / 1e3] * len(distances)   # dB/m, same for every link
        total = sum(distances)
        for name, algo in case["algorithms"].items():
            alloc = algo["alloc"]
            pairs = link_pairs(alloc)
            if len(pairs) != len(distances):
                raise ValueError(f"{case['sheet']} case {case['case']} {name}: allocation covers {len(pairs)} "
                                 f"links but p_gen has {len(distances)}")
            jobs.append({
                "sheet": case["sheet"], "case": case["case"], "algorithm": name,
                "order_sheet": algo["order_sheet"],
                "order_sim": [x - 1 for x in algo["order_sheet"]],   # 1-indexed -> 0-indexed
                "distances_km": distances, "total_km": total, "attenuations": attenuations,
                "p_gen": case["p_gen"], "p_swap": case["p_swap"],
                "alloc_pairs": pairs,
                "capacity": max(left + right for left, right in alloc.values()),
                "expected": algo["expected"],
                "reference": "; ".join(f"{k}={v}" for k, v in case["reference"].items()),
                "seed": seed, "generation": generation,
            })
    return jobs


# ----------------------------------------------------------------------------------------
# Running
# ----------------------------------------------------------------------------------------
def run_job(job, coherence_ms, duration_s, target_e2e):
    es.SIM_DURATION_S = duration_s
    es.TARGET_E2E = target_e2e
    tick = time.time()
    hardware = {}
    if job["generation"] == "single_heralded":
        hardware = {"memory_efficiency": SH_MEMORY_EFFICIENCY, "detector_efficiency": SH_DETECTOR_EFFICIENCY,
                    "bsm_success_rate": SH_BSM_SUCCESS_RATE}
    r = es.run_once(job["order_sim"], job["distances_km"], job["alloc_pairs"], coherence_ms,
                    seed=job["seed"], swap_prob=job["p_swap"], memory_capacity=job["capacity"],
                    generation=job["generation"], attenuations=job["attenuations"], **hardware)
    return job, r, time.time() - tick


def safe_run_job(job, coherence_ms, duration_s, target_e2e):
    """run_job that returns (job, None, traceback) instead of raising, so one bad run can't stop a batch."""
    try:
        return run_job(job, coherence_ms, duration_s, target_e2e) + (None,)
    except Exception:
        return job, None, 0.0, traceback.format_exc()


def record_failure(errors_log, job, tb):
    with open(errors_log, "a", encoding="utf-8") as f:
        f.write(f"===== {time.strftime('%Y-%m-%d %H:%M:%S')}  {describe(job)}\n{tb}\n")


def job_key(sheet, case, algorithm, seed):
    return (sheet, str(case), algorithm, str(seed))


def load_done(results_csv):
    if not os.path.exists(results_csv):
        return set()
    with open(results_csv, newline="", encoding="utf-8") as f:
        return {job_key(r["sheet"], r["case"], r["algorithm"], r["seed"]) for r in csv.DictReader(f)}


def write_result(results_csv, job, r, wall):
    new_file = not os.path.exists(results_csv)
    with open(results_csv, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow({
            "sheet": job["sheet"], "case": job["case"], "algorithm": job["algorithm"],
            "order_sheet": " ".join(map(str, job["order_sheet"])),
            "order_sim": " ".join(map(str, job["order_sim"])),
            "distances_km": " ".join(f"{d:.3f}" for d in job["distances_km"]),
            "total_km": f"{job['total_km']:.3f}",
            "attenuation_db_km": " ".join(f"{a * 1e3:.4f}" for a in job["attenuations"]),
            "p_gen": " ".join(map(str, job["p_gen"])), "p_swap": " ".join(map(str, job["p_swap"])),
            "alloc_pairs": " ".join(map(str, job["alloc_pairs"])),
            "expected_ent_sheet": job["expected"], "reference_expected_ent_sheet": job["reference"],
            "generation": job["generation"], "seed": job["seed"], "e2e": r["e2e"], "elapsed_s": f"{r['elapsed_s']:.6f}",
            "throughput_ent_s": f"{r['throughput']:.4f}", "expired": r["expired"], "wall_s": f"{wall:.1f}",
        })


def describe(job):
    return (f"{job['sheet']:<22} case {job['case']:>3}  {job['algorithm']:<4} "
            f"order {job['order_sheet']} -> {job['order_sim']}  "
            f"L = [{', '.join(f'{d:.1f}' for d in job['distances_km'])}] km ({job['total_km']:.1f})  "
            f"loss = [{', '.join(f'{a * 1e3:.3f}' for a in job['attenuations'])}] dB/km  "
            f"q = {job['p_swap']}  pairs = {job['alloc_pairs']}")


def summarize(results_csv, sheets):
    if not os.path.exists(results_csv):
        return
    with open(results_csv, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    print("\n" + "=" * 78)
    print(f"{'sheet':<24}{'algorithm':<10}{'runs':>6}{'mean ent/s':>14}{'best in case':>16}")
    for sheet in sheets:
        sheet_rows = [r for r in rows if r["sheet"] == sheet]
        if not sheet_rows:
            continue
        by_case = {}
        for r in sheet_rows:
            by_case.setdefault(r["case"], []).append(r)
        names = list(dict.fromkeys(r["algorithm"] for r in sheet_rows))
        wins = {name: 0 for name in names}
        for case_rows in by_case.values():
            best = max(case_rows, key=lambda r: float(r["throughput_ent_s"]))
            wins[best["algorithm"]] += 1
        for name in names:
            alg_rows = [r for r in sheet_rows if r["algorithm"] == name]
            mean = sum(float(r["throughput_ent_s"]) for r in alg_rows) / len(alg_rows)
            print(f"{sheet:<24}{name:<10}{len(alg_rows):>6}{mean:>14.2f}{wins[name]:>16}")
    print("=" * 78)


def main():
    parser = argparse.ArgumentParser(description="Simulate path-graph test cases from an Excel workbook in SeQUeNCe")
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK, help="Excel file with the test cases")
    parser.add_argument("--sheets", nargs="+", default=None, help="sheet names to run (default: all sheets)")
    parser.add_argument("--out", default=None, help="results CSV (default: <workbook>_results_<generation>.csv)")
    parser.add_argument("--generation", choices=["single_heralded", "barret_kok"], default=GENERATION,
                        help="link entanglement generation (single_heralded: attempt success == p_gen exactly)")
    parser.add_argument("--cases", type=int, default=CASES_PER_SHEET, help="cases per sheet (first N)")
    parser.add_argument("--coherence", type=float, default=es.COHERENCE_TIME_MS, help="coherence time in ms")
    parser.add_argument("--duration", type=float, default=es.SIM_DURATION_S, help="max simulated seconds per run")
    parser.add_argument("--target", type=int, default=es.TARGET_E2E, help="stop a run after this many E2E")
    parser.add_argument("--seed", type=int, default=0)
    parser.add_argument("--jobs", type=int, default=1, help="runs in parallel")
    parser.add_argument("--dry-run", action="store_true", help="print the mapped cases without simulating")
    args = parser.parse_args()

    workbook = os.path.abspath(args.workbook)
    available = openpyxl.load_workbook(workbook, read_only=True).sheetnames
    sheets = args.sheets if args.sheets else available
    for sheet in sheets:
        if sheet not in available:
            raise ValueError(f"unknown sheet {sheet!r}; {os.path.basename(workbook)} has {available}")
    results_csv = (os.path.abspath(args.out) if args.out
                   else f"{os.path.splitext(workbook)[0]}_results_{args.generation}.csv")

    jobs = build_jobs(read_cases(workbook, sheets, args.cases), args.seed, args.generation)
    if not jobs:
        raise ValueError("no simulatable rows found (need an Order with a node tuple and a Qubit Allocation)")
    done = load_done(results_csv)
    todo = [j for j in jobs if job_key(j["sheet"], j["case"], j["algorithm"], j["seed"]) not in done]

    algorithms = list(dict.fromkeys(j["algorithm"] for j in jobs))
    print(f"Workbook: {workbook}")
    print(f"Sheets:   {sheets}")
    print(f"Results:  {results_csv}")
    print(f"{len(jobs)} runs (up to {args.cases} cases per sheet; algorithms {algorithms}), "
          f"{len(jobs) - len(todo)} already done, {len(todo)} to run")
    print(f"T_cohere = {args.coherence:g} ms, each run: {args.duration:g} s or {args.target} E2E, seed {args.seed}")
    if args.generation == "single_heralded":
        print(f"Generation: single-heralded, BSM success {SH_BSM_SUCCESS_RATE:g}, memory eff {SH_MEMORY_EFFICIENCY:g}, "
              f"detector eff {SH_DETECTOR_EFFICIENCY:g}  ->  attempt success = p_gen")
    else:
        print("Generation: Barrett-Kok with Section V-A efficiencies  ->  attempt success ~0.41 x p_gen")
    print(f"Longest path: {max(j['total_km'] for j in jobs):.1f} km\n")

    if args.dry_run:
        for job in jobs:
            print(describe(job))
        return

    errors_log = os.path.splitext(results_csv)[0] + "_errors.log"
    failures = []

    def handle(i, job, r, wall, tb):
        if tb is not None:
            failures.append(job)
            record_failure(errors_log, job, tb)
            print(f"[{i}/{len(todo)}] FAILED {describe(job)}\n    -> {tb.strip().splitlines()[-1]} "
                  f"(traceback in {os.path.basename(errors_log)})", flush=True)
            return
        write_result(results_csv, job, r, wall)
        print(f"[{i}/{len(todo)}] {describe(job)}\n    -> {r['e2e']} E2E in {r['elapsed_s']:.3f} s "
              f"= {r['throughput']:.2f} ent/s (wall {wall:.1f} s)", flush=True)

    if args.jobs <= 1:
        for i, job in enumerate(todo, 1):
            handle(i, *safe_run_job(job, args.coherence, args.duration, args.target))
    else:
        with ProcessPoolExecutor(max_workers=args.jobs) as pool:
            futures = {pool.submit(safe_run_job, job, args.coherence, args.duration, args.target): job
                       for job in todo}
            for i, future in enumerate(as_completed(futures), 1):
                try:
                    result = future.result()
                except Exception:   # the worker process itself died (e.g. killed, out of memory)
                    result = (futures[future], None, 0.0, traceback.format_exc())
                handle(i, *result)

    summarize(results_csv, sheets)
    if failures:
        print(f"\n{len(failures)} run(s) failed; tracebacks are in {errors_log}. "
              f"Failed runs are not in the CSV, so running the same command again retries them.")


if __name__ == "__main__":
    main()
