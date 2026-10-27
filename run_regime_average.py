"""Average throughput per algorithm in each regime (sheet), over repeated runs and/or all cases.

Two modes, for each sheet of the workbook (Wide Random, High Fidelity, Mid-Noise, Random Bottleneck 30%):

  Default (one case):  one test case is chosen at random (fixed --pick-seed, so the same cases are
                       chosen on rerun; or set explicitly with --case "Wide Random=3") and every
                       algorithm is run --runs times (default 30) on it.

  --all-cases:         every test case in the sheet is used, each algorithm run --runs times per
                       case (default 1). Results are averaged per case first, then across cases.

Every algorithm with a swap order and allocation in a case (DP, Our Algo, CBT, ...) is simulated
with the same mapping as run_5node_cases.py; each run lasts --duration s (20) of simulated time or
until --target E2E entanglements (200).

Metrics:
  * Throughput (ent/s): E2E entanglements per second of simulated time.
  * Expired Memories per E2E: memories that decohered before use / E2E delivered, per run
    (runs with no E2E are left out of this average).
  * --all-cases only:
      - Throughput vs CBT: per case, algorithm throughput / CBT throughput, averaged over cases
        (a fair cross-case comparison, since raw throughput differs hugely between cases).
      - Cases Won: number of cases where the algorithm had the highest throughput.

Output: an Excel file with an "Averages" sheet (one row per regime and algorithm) and, with
--all-cases, a "Per case" sheet (one row per case and algorithm, with order and allocation).
Every finished run is appended to a CSV first, so an interrupted batch resumes where it stopped.

Usage:
    python local/run_regime_average.py --jobs 20                          # 1 random case x 30 runs
    python local/run_regime_average.py --all-cases --jobs 20              # 10 cases x 1 run
    python local/run_regime_average.py --all-cases --runs 3 --jobs 20     # 10 cases x 3 runs
    python local/run_regime_average.py --case "Wide Random=3" --jobs 20
    python local/run_regime_average.py --all-cases --dry-run              # show the cases only
    python local/run_regime_average.py --all-cases --runs 10 --jobs 20     #to run
"""

import argparse
import csv
import os
import random
import statistics
import sys
import traceback
from concurrent.futures import ProcessPoolExecutor, as_completed

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import run_5node_cases as rc

RUNS_ONE_CASE = 30
RUNS_ALL_CASES = 1
DURATION_S = 20
TARGET_E2E = 200
PICK_SEED = 2026
BASELINE = "cbt"                     # algorithm used for "Throughput vs CBT"

ALGORITHM_NAMES = {"dp": "DP", "ours": "Our Algo", "cbt": "Baseline (CBT, Equal Alloc)", "gso": "(GSO, EA)"}
RUN_FIELDS = ["regime", "case", "algorithm", "run", "order", "allocation", "e2e", "elapsed_s",
              "throughput_ent_s", "expired", "expired_per_e2e"]


def format_allocation(alloc):
    return "{" + ", ".join(f"{node}: ({left}, {right})" for node, (left, right) in sorted(alloc.items())) + "}"


def display_name(name):
    return ALGORITHM_NAMES.get(name, name)


# ----------------------------------------------------------------------------------------
# Choosing cases and building jobs
# ----------------------------------------------------------------------------------------
def choose_cases(cases, sheets, explicit, pick_seed, all_cases):
    """{sheet: [case, ...]}: every case with all_cases, else one random (or explicit) case."""
    rng = random.Random(pick_seed)
    chosen = {}
    for sheet in sheets:
        sheet_cases = [c for c in cases if c["sheet"] == sheet]
        if not sheet_cases:
            raise ValueError(f"sheet {sheet!r} has no test cases")
        if all_cases:
            chosen[sheet] = sheet_cases
            continue
        pick = rng.choice(sheet_cases)            # always draw, so other sheets' picks stay stable
        if sheet in explicit:
            matches = [c for c in sheet_cases if c["case"] == explicit[sheet]]
            if not matches:
                raise ValueError(f"sheet {sheet!r} has no test case {explicit[sheet]}")
            pick = matches[0]
        chosen[sheet] = [pick]
    return chosen


def build_run_jobs(chosen, runs, generation):
    jobs = []
    for sheet_cases in chosen.values():
        for case in sheet_cases:
            for run in range(runs):
                for job in rc.build_jobs([case], run, generation):
                    job["allocation_text"] = format_allocation(case["algorithms"][job["algorithm"]]["alloc"])
                    jobs.append(job)
    return jobs


# ----------------------------------------------------------------------------------------
# Per-run CSV
# ----------------------------------------------------------------------------------------
def load_done(path):
    if not os.path.exists(path):
        return {}
    with open(path, newline="", encoding="utf-8") as f:
        return {(r["regime"], r["case"], r["algorithm"], r["run"]): r for r in csv.DictReader(f)}


def append_run(path, job, r):
    new_file = not os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=RUN_FIELDS)
        if new_file:
            writer.writeheader()
        writer.writerow({
            "regime": job["sheet"], "case": job["case"], "algorithm": job["algorithm"], "run": job["seed"],
            "order": str(tuple(job["order_sheet"])), "allocation": job["allocation_text"],
            "e2e": r["e2e"], "elapsed_s": f"{r['elapsed_s']:.6f}", "throughput_ent_s": f"{r['throughput']:.4f}",
            "expired": r["expired"], "expired_per_e2e": f"{r['expired'] / r['e2e']:.4f}" if r["e2e"] else "",
        })


# ----------------------------------------------------------------------------------------
# Summaries
# ----------------------------------------------------------------------------------------
def mean_or_none(values):
    return statistics.mean(values) if values else None


def per_case_results(rows, chosen, runs_wanted):
    """One row per (regime, case, algorithm): means over that case's runs."""
    results = []
    for sheet, sheet_cases in chosen.items():
        for case in sheet_cases:
            for name in case["algorithms"]:
                runs = [r for r in rows if r["regime"] == sheet and r["case"] == str(case["case"])
                        and r["algorithm"] == name]
                if not runs:
                    continue
                results.append({
                    "Regime": sheet, "Test Case": case["case"], "Algorithm": display_name(name), "_name": name,
                    "Swap Order": runs[0]["order"], "Qubit Allocation": runs[0]["allocation"],
                    "Throughput (ent/s)": statistics.mean(float(r["throughput_ent_s"]) for r in runs),
                    "Expired Memories per E2E": mean_or_none(
                        [float(r["expired_per_e2e"]) for r in runs if r["expired_per_e2e"]]),
                    "Runs": len(runs), "_complete": len(runs) >= runs_wanted,
                })
    return results


def regime_averages(case_rows, chosen):
    """One row per (regime, algorithm): averages over the regime's cases."""
    summary = []
    for sheet, sheet_cases in chosen.items():
        rows = [r for r in case_rows if r["Regime"] == sheet]
        names = list(dict.fromkeys(r["_name"] for r in rows))
        by_case = {}
        for r in rows:
            by_case.setdefault(r["Test Case"], {})[r["_name"]] = r
        wins = {name: 0 for name in names}
        for algs in by_case.values():
            if len(algs) == len(names):          # only compare cases where every algorithm finished
                wins[max(algs, key=lambda n: algs[n]["Throughput (ent/s)"])] += 1
        for name in names:
            alg_rows = [r for r in rows if r["_name"] == name]
            ratios = [r["Throughput (ent/s)"] / by_case[r["Test Case"]][BASELINE]["Throughput (ent/s)"]
                      for r in alg_rows
                      if BASELINE in by_case[r["Test Case"]]
                      and by_case[r["Test Case"]][BASELINE]["Throughput (ent/s)"] > 0]
            orders = {r["Swap Order"] for r in alg_rows}
            allocations = {r["Qubit Allocation"] for r in alg_rows}
            summary.append({
                "Regime": sheet, "Algorithm": display_name(name), "Cases": len(alg_rows),
                "Test Case": alg_rows[0]["Test Case"] if len(alg_rows) == 1 else f"{len(alg_rows)} cases",
                "Swap Order": orders.pop() if len(orders) == 1 else "varies (see Per case)",
                "Qubit Allocation": allocations.pop() if len(allocations) == 1 else "varies (see Per case)",
                "Throughput (ent/s)": statistics.mean(r["Throughput (ent/s)"] for r in alg_rows),
                "Expired Memories per E2E": mean_or_none(
                    [r["Expired Memories per E2E"] for r in alg_rows if r["Expired Memories per E2E"] is not None]),
                "Throughput vs CBT": mean_or_none(ratios),
                "Cases Won": wins[name],
                "Runs": sum(r["Runs"] for r in alg_rows),
            })
    return summary


# ----------------------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------------------
def write_sheet(ws, columns, rows, widths, number_formats):
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2A78D6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([row[c] for c in columns])
    for row in ws.iter_rows(min_row=2):
        for cell, column in zip(row, columns):
            if column in number_formats:
                cell.number_format = number_formats[column]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def write_excel(path, summary, case_rows, all_cases, runs_wanted, duration, target):
    fmt = {"Throughput (ent/s)": "#,##0.00", "Expired Memories per E2E": "#,##0.00", "Throughput vs CBT": "0.00x"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Averages"
    if all_cases:
        columns = ["Regime", "Algorithm", "Cases", "Throughput (ent/s)", "Expired Memories per E2E",
                   "Throughput vs CBT", "Cases Won"]
        write_sheet(ws, columns, summary, [24, 28, 8, 18, 22, 16, 11], fmt)
        note = (f"Averages over the regime's test cases; each case value is the mean of {runs_wanted} run(s), "
                f"each lasting {duration:g} s of simulated time or until {target} E2E. Throughput vs CBT is the "
                f"per-case ratio to CBT, averaged; Cases Won counts cases with the highest throughput.")
        detail = wb.create_sheet("Per case")
        detail_columns = ["Regime", "Test Case", "Algorithm", "Swap Order", "Qubit Allocation",
                          "Throughput (ent/s)", "Expired Memories per E2E", "Runs"]
        write_sheet(detail, detail_columns, case_rows, [24, 10, 28, 12, 48, 18, 22, 7], fmt)
    else:
        columns = ["Regime", "Test Case", "Algorithm", "Swap Order", "Qubit Allocation",
                   "Throughput (ent/s)", "Expired Memories per E2E"]
        write_sheet(ws, columns, summary, [24, 10, 28, 12, 48, 18, 22], fmt)
        note = (f"Each value is the mean over {runs_wanted} runs of one test case; each run lasted {duration:g} s "
                f"of simulated time or until {target} E2E entanglements.")
    note_row = ws.max_row + 2
    ws.cell(note_row, 1, note)
    incomplete = [f"{r['Regime']} case {r['Test Case']} / {r['Algorithm']} ({r['Runs']} of {runs_wanted} runs)"
                  for r in case_rows if not r["_complete"]]
    if incomplete:
        ws.cell(note_row + 1, 1, "Incomplete: " + "; ".join(incomplete))
    wb.save(path)


def print_summary(summary, all_cases):
    print("\n" + "=" * 110)
    if all_cases:
        print(f"{'Regime':<24}{'Algorithm':<30}{'Cases':>6}{'Throughput':>14}{'Expired/E2E':>14}"
              f"{'vs CBT':>9}{'Won':>6}")
        for s in summary:
            exp = f"{s['Expired Memories per E2E']:.2f}" if s["Expired Memories per E2E"] is not None else "-"
            ratio = f"{s['Throughput vs CBT']:.2f}x" if s["Throughput vs CBT"] is not None else "-"
            print(f"{s['Regime']:<24}{s['Algorithm']:<30}{s['Cases']:>6}{s['Throughput (ent/s)']:>14.2f}"
                  f"{exp:>14}{ratio:>9}{s['Cases Won']:>6}")
    else:
        print(f"{'Regime':<24}{'Case':>5}  {'Algorithm':<30}{'Order':<12}{'Throughput':>14}{'Expired/E2E':>14}"
              f"{'Runs':>6}")
        for s in summary:
            exp = f"{s['Expired Memories per E2E']:.2f}" if s["Expired Memories per E2E"] is not None else "-"
            print(f"{s['Regime']:<24}{s['Test Case']:>5}  {s['Algorithm']:<30}{s['Swap Order']:<12}"
                  f"{s['Throughput (ent/s)']:>14.2f}{exp:>14}{s['Runs']:>6}")
    print("=" * 110)


# ----------------------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="Average throughput per algorithm in each regime")
    parser.add_argument("--workbook", default=rc.DEFAULT_WORKBOOK)
    parser.add_argument("--sheets", nargs="+", default=None, help="regimes (sheets) to run (default: all)")
    parser.add_argument("--all-cases", action="store_true", help="use every test case in each sheet")
    parser.add_argument("--case", action="append", default=[], metavar='"SHEET=N"',
                        help="one-case mode: use test case N for SHEET instead of a random one (repeatable)")
    parser.add_argument("--pick-seed", type=int, default=PICK_SEED, help="seed for the random case choice")
    parser.add_argument("--runs", type=int, default=None,
                        help=f"runs per case and algorithm (default {RUNS_ONE_CASE}, or {RUNS_ALL_CASES} with --all-cases)")
    parser.add_argument("--duration", type=float, default=DURATION_S, help="max simulated seconds per run")
    parser.add_argument("--target", type=int, default=TARGET_E2E, help="stop a run after this many E2E")
    parser.add_argument("--coherence", type=float, default=rc.es.COHERENCE_TIME_MS, help="coherence time in ms")
    parser.add_argument("--generation", choices=["single_heralded", "barret_kok"], default=rc.GENERATION)
    parser.add_argument("--jobs", type=int, default=1, help="runs in parallel")
    parser.add_argument("--out", default=None, help="Excel summary file")
    parser.add_argument("--runs-csv", default=None, help="per-run CSV file")
    parser.add_argument("--dry-run", action="store_true", help="only show the chosen cases")
    args = parser.parse_args()

    if args.all_cases and args.case:
        raise ValueError("--case picks one case per sheet; it cannot be combined with --all-cases")
    runs_wanted = args.runs if args.runs is not None else (RUNS_ALL_CASES if args.all_cases else RUNS_ONE_CASE)

    workbook = os.path.abspath(args.workbook)
    # separate files per mode and per generation protocol, so their runs never mix on resume
    # (single-heralded keeps the original names)
    stem = (os.path.splitext(workbook)[0] + ("_allcases" if args.all_cases else "")
            + ("" if args.generation == "single_heralded" else f"_{args.generation}"))
    out_xlsx = os.path.abspath(args.out) if args.out else stem + "_regime_averages.xlsx"
    runs_csv = os.path.abspath(args.runs_csv) if args.runs_csv else stem + "_regime_runs.csv"
    errors_log = os.path.splitext(runs_csv)[0] + "_errors.log"

    available = openpyxl.load_workbook(workbook, read_only=True).sheetnames
    sheets = args.sheets or available
    explicit = {}
    for item in args.case:
        sheet, _, number = item.rpartition("=")
        explicit[sheet] = int(number)
    for sheet in list(sheets) + list(explicit):
        if sheet not in available:
            raise ValueError(f"unknown sheet {sheet!r}; {os.path.basename(workbook)} has {available}")

    cases = rc.read_cases(workbook, sheets, cases_per_sheet=10 ** 9)
    chosen = choose_cases(cases, sheets, explicit, args.pick_seed, args.all_cases)

    print(f"Workbook: {workbook}")
    print(f"Mode: {'all cases' if args.all_cases else 'one case per regime'}; {runs_wanted} run(s) per case and "
          f"algorithm; each run {args.duration:g} s or {args.target} E2E; generation {args.generation}; "
          f"T_cohere {args.coherence:g} ms")
    for sheet, sheet_cases in chosen.items():
        if args.all_cases:
            print(f"\n{sheet}: test cases {[c['case'] for c in sheet_cases]}")
            continue
        case = sheet_cases[0]
        print(f"\n{sheet}: test case {case['case']}" + ("  (chosen with --case)" if sheet in explicit else "  (random)"))
        km = [rc.pgen_to_km(p) for p in case["p_gen"]]
        print(f"    p_gen {case['p_gen']} -> L = [{', '.join(f'{d:.1f}' for d in km)}] km,  p_swap {case['p_swap']}")
        for name, algo in case["algorithms"].items():
            print(f"    {display_name(name):<28} order {algo['order_sheet']}  alloc {format_allocation(algo['alloc'])}")
    if args.dry_run:
        return

    jobs = build_run_jobs(chosen, runs_wanted, args.generation)
    done = load_done(runs_csv)
    todo = [j for j in jobs if (j["sheet"], str(j["case"]), j["algorithm"], str(j["seed"])) not in done]
    print(f"\n{len(jobs)} runs in total, {len(jobs) - len(todo)} already in {os.path.basename(runs_csv)}, "
          f"{len(todo)} to run\n")

    failures = 0

    def handle(i, job, r, wall, tb):
        nonlocal failures
        label = f"[{i}/{len(todo)}] {job['sheet']} case {job['case']} {job['algorithm']} run {job['seed']}"
        if tb is not None:
            failures += 1
            rc.record_failure(errors_log, job, tb)
            print(f"{label}: FAILED ({tb.strip().splitlines()[-1]})", flush=True)
            return
        append_run(runs_csv, job, r)
        print(f"{label}: {r['e2e']} E2E in {r['elapsed_s']:.3f} s = {r['throughput']:.2f} ent/s, "
              f"{r['expired']} expired (wall {wall:.0f} s)", flush=True)

    if args.jobs <= 1:
        for i, job in enumerate(todo, 1):
            handle(i, *rc.safe_run_job(job, args.coherence, args.duration, args.target))
    else:
        with ProcessPoolExecutor(max_workers=args.jobs) as pool:
            futures = {pool.submit(rc.safe_run_job, job, args.coherence, args.duration, args.target): job
                       for job in todo}
            for i, future in enumerate(as_completed(futures), 1):
                try:
                    result = future.result()
                except Exception:   # the worker process itself died
                    result = (futures[future], None, 0.0, traceback.format_exc())
                handle(i, *result)

    rows = list(load_done(runs_csv).values())
    case_rows = per_case_results(rows, chosen, runs_wanted)
    summary = regime_averages(case_rows, chosen) if args.all_cases else case_rows
    write_excel(out_xlsx, summary, case_rows, args.all_cases, runs_wanted, args.duration, args.target)
    print_summary(summary, args.all_cases)
    print(f"Saved {out_xlsx}")
    if failures:
        print(f"{failures} run(s) failed (see {errors_log}); run the same command again to retry them.")


if __name__ == "__main__":
    main()
